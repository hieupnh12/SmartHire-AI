"""Fill docs/test/Report5_Unit Test.xls from the real JUnit tests (Windows + Microsoft Excel required).

One sheet per function under test (class + method). Each @Test method becomes one UTCID column:
  - condition / confirmation rows are derived from the test source (scenario, exception, HTTP status),
  - Type N/A/B is inferred from the test name and assertions,
  - Passed/Failed comes from backend/target/surefire-reports (run `mvnw test jacoco:report` first),
  - Executed Date is the date the test was last changed (git blame; CI runs mvn test on every push),
    or the surefire run date for failing tests,
  - Created By is the git author of the test.
Re-runnable: function sheets are regenerated, and a new Record of change version is added to the Cover
only when test cases or results changed. Optional overrides: scripts/testing/content/unit_test/<Class>.json

Usage:
  python scripts/testing/fill_unit_test.py [--dry-run] [--note "..."] [--report path.xls]
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_traceability import MAIN, SUBJECT_SUFFIXES, TEST, guess_kind, norm_path, parse_controllers  # noqa: E402
from fill_test_report import HISTORY_KEYS, load_content, next_version  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "docs" / "test" / "Report5_Unit Test.xls"
CONTENT_DIR = Path(__file__).resolve().parent / "content" / "unit_test"
SUREFIRE = ROOT / "backend" / "target" / "surefire-reports"
JACOCO_CSV = ROOT / "backend" / "target" / "site" / "jacoco" / "jacoco.csv"
MATRIX = ROOT / ".tmp" / "testing" / "Traceability_Matrix.xlsx"
BACKUP_DIR = ROOT / ".tmp" / "testing" / "backup"

TEMPLATE_SHEET = "_Template"
SNAPSHOT_SHEET = "_Snapshot"
SAMPLE_SHEETS = ("Function 1", "Function 2", "Function3", "Example")
FIXED_SHEETS = ("Guideline", "Cover", "Functions", "Statistics")

AUTHORS = {
    "trexloc": "Lê Mỹ Lộc",
    "hieupnh12": "Nguyễn Minh Hiểu",
    "minhieu.212": "Nguyễn Minh Hiểu",
    "natto-sin": "Nguyễn Nhất Sinh",
    "CuongLC-1301": "Lê Công Cường",
    "tienlp204": "Lê Phúc Tiền",
}
EXECUTED_BY = "Automated: mvn test (GitHub Actions CI)"
ENVIRONMENT = (
    "1. Java 21, Maven 3 (mvn test), JUnit 5, Mockito, Spring Boot Test, AssertJ.\n"
    "2. Code coverage: JaCoCo 0.8.12 (backend/target/site/jacoco).\n"
    "3. Integration tests: H2 / Testcontainers-style setup defined in each test class.\n"
    "4. CI: GitHub Actions workflow ci.yml runs mvn -B test on every push and pull request."
)
HTTP_STATUS = {"isOk": "HTTP 200", "isCreated": "HTTP 201", "isAccepted": "HTTP 202", "isNoContent": "HTTP 204",
               "is2xxSuccessful": "HTTP 2xx", "isBadRequest": "HTTP 400", "isUnauthorized": "HTTP 401",
               "isForbidden": "HTTP 403", "isNotFound": "HTTP 404", "isConflict": "HTTP 409",
               "isUnprocessableEntity": "HTTP 422", "is4xxClientError": "HTTP 4xx",
               "isInternalServerError": "HTTP 500", "isServiceUnavailable": "HTTP 503"}
XL_SHIFT_DOWN, XL_SHIFT_RIGHT, XL_PASTE_FORMATS = -4121, -4161, -4122


# ------------------------------------------------------------------ model --
@dataclass
class UnitCase:
    test_class: str
    method: str
    scenario: str
    kind: str
    exception: str
    returns: str
    status: str  # P / F / "" (untested)
    executed: datetime | None
    author: str


@dataclass
class Function:
    cls: str
    name: str
    precondition: str
    loc: int | None
    cases: list[UnitCase] = field(default_factory=list)
    features: set[str] = field(default_factory=set)
    sheet: str = ""
    code: str = ""


def humanize(name: str) -> str:
    words = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", name.replace("_", " ")).split()
    return " ".join(words).capitalize()


def blame(path: Path) -> list[tuple[str, datetime]]:
    out = subprocess.run(["git", "blame", "--line-porcelain", "--", str(path.relative_to(ROOT))], cwd=ROOT,
                         capture_output=True, text=True, encoding="utf-8")
    lines, author, when = [], "", datetime.now()
    for line in out.stdout.splitlines():
        if line.startswith("author "):
            author = line[7:]
        elif line.startswith("author-time "):
            when = datetime.fromtimestamp(int(line[12:]))
        elif line.startswith("\t"):
            lines.append((author, when))
    return lines


def load_results() -> tuple[dict[str, dict[str, str]], dict[str, datetime]]:
    """class -> {test method -> P/F/skip}, class -> run time."""
    results, runs = defaultdict(dict), {}
    for path in SUREFIRE.glob("TEST-*.xml"):
        root = ET.parse(path).getroot()
        cls = root.get("name").split(".")[-1]
        stamp = root.get("timestamp")
        runs[cls] = datetime.fromisoformat(stamp) if stamp else datetime.fromtimestamp(path.stat().st_mtime)
        for tc in root.iter("testcase"):
            name = re.split(r"[(\[]", tc.get("name"))[0]
            state = ("F" if tc.find("failure") is not None or tc.find("error") is not None
                     else "skip" if tc.find("skipped") is not None else "P")
            if results[cls].get(name) != "F":
                results[cls][name] = state
    return results, runs


def method_loc(path: Path, name: str) -> int | None:
    src = path.read_text(encoding="utf-8")
    m = re.search(r"\n[ \t]+(?:public|protected|private)?[\w<>\[\],.?\s]*?\b" + re.escape(name) + r"\s*\([^;{]*?\)"
                  r"[^;{]*\{", src)
    if not m:
        return None
    depth, i = 1, m.end()
    while depth and i < len(src):
        depth += {"{": 1, "}": -1}.get(src[i], 0)
        i += 1
    return src[m.start():i].strip().count("\n") + 1


def precondition_of(src: str) -> str:
    if "@WebMvcTest" in src:
        return "Spring MVC test slice; services are mocked (@MockBean)"
    if "@SpringBootTest" in src or "@DataJpaTest" in src:
        return "Spring application context with test database"
    if "Mockito" in src or "@Mock" in src:
        return "Dependencies are mocked with Mockito"
    return "Object created directly with test data"


def test_bodies(src: str):
    for m in re.finditer(r"@(?:Test|ParameterizedTest)\b", src):
        decl = re.search(r"void\s+(\w+)\s*\(", src[m.end():])
        if not decl:
            continue
        start = m.end() + decl.start()
        end_m = re.search(r"\n    (?:@\w|(?:public|private|protected|static|void)\b)", src[start + 1:])
        end = start + 1 + end_m.start() if end_m else len(src)
        display = re.search(r'@DisplayName\s*\(\s*"([^"]*)"', src[m.start():start])
        yield decl.group(1), (display.group(1) if display else ""), src[start:end], src[:start].count("\n")


def analyse(classes: dict[str, Path], test_features: dict[str, set[str]]) -> list[Function]:
    endpoints, _ = parse_controllers(classes)
    ep_index = {(e.method, e.path): e for e in endpoints}
    results, runs = load_results()
    functions: dict[tuple[str, str], Function] = {}

    for path in sorted(TEST.rglob("*.java")):
        src = path.read_text(encoding="utf-8")
        test_class = path.stem
        base = re.sub(r"(Integration)?Tests?$", "", test_class)
        subjects: dict[str, str] = {}
        for m in re.finditer(r"@(?:InjectMocks|Autowired)\s+(?:private\s+)?(?:final\s+)?(\w+)\s+(\w+)", src):
            subjects[m.group(2)] = m.group(1)
        for m in re.finditer(r"^[ \t]+(?:(?:private|protected|final)\s+)*(\w+)\s+(\w+)\s*[;=]", src, re.M):
            if m.group(1) in classes and m.group(1).endswith(SUBJECT_SUFFIXES):
                subjects.setdefault(m.group(2), m.group(1))
        for m in re.finditer(r"\b(\w+)\s+(\w+)\s*=\s*new\s+(\w+)\s*\(", src):
            if m.group(3) in classes and m.group(3).endswith(SUBJECT_SUFFIXES):
                subjects[m.group(2)] = m.group(3)
        subjects = {v: t for v, t in subjects.items() if t in classes}
        preferred = sorted(subjects, key=lambda v: subjects[v] != base)
        webmvc = re.search(r"@WebMvcTest\s*\(\s*(?:controllers\s*=\s*)?\{?\s*(\w+)\.class", src)
        precondition = precondition_of(src)
        lines = blame(path)

        for method, display, body, line_no in test_bodies(src):
            cls, name = (base if base in classes else test_class), "multipleMethods"
            for var in preferred:
                call = re.search(r"\b" + re.escape(var) + r"\.(\w+)\s*\(", body)
                if call:
                    cls, name = subjects[var], call.group(1)
                    break
            else:
                perform = re.search(r"perform\(\s*(get|post|put|patch|delete)\(\s*\"([^\"]+)\"", body)
                if perform:
                    http = perform.group(1).upper()
                    route = norm_path(re.sub(r"/\d+(?=/|$)", "/{}", perform.group(2)))
                    ep = ep_index.get((http, route))
                    cls = ep.controller if ep else (webmvc.group(1) if webmvc else cls)
                    name = ep.handler if ep else f"{http} {route}"

            exc = re.search(r"(?:assertThrows|isInstanceOf|assertThatExceptionOfType)\(\s*(\w+)\.class", body)
            code = re.search(r"\"([A-Z][A-Z0-9]*(?:_[A-Z0-9]+)+)\"", body)
            exception = f"{exc.group(1)}{f' ({code.group(1)})' if code else ''}" if exc else ""
            http_status = re.search(r"status\(\)\.(is\w+)\(\)", body)
            returns = (HTTP_STATUS.get(http_status.group(1), http_status.group(1)) if http_status
                       else "" if exception else "Returns expected result")
            kind = "A" if exception or re.match(r"HTTP [45]", returns) else guess_kind(method + " " + display)

            state = results.get(test_class, {}).get(method, "")
            body_lines = lines[line_no: line_no + body.count("\n") + 1] or [("", datetime.now())]
            if state == "F":
                executed = runs.get(test_class)
            elif state == "P":
                executed = max(when for _, when in body_lines)
            else:
                executed = None
            author = Counter(a for a, _ in body_lines).most_common(1)[0][0]

            fn = functions.get((cls, name))
            if not fn:
                fn = functions[(cls, name)] = Function(cls, name, precondition,
                                                       method_loc(classes[cls], name) if cls in classes else None)
            fn.features |= test_features.get(test_class, set())
            fn.cases.append(UnitCase(test_class, method, display or humanize(method), kind, exception, returns,
                                     "" if state == "skip" else state, executed, AUTHORS.get(author, author)))

    ordered = sorted(functions.values(), key=lambda f: (f.cls, f.name))
    used: set[str] = set()
    for i, fn in enumerate(ordered, 1):
        fn.code = f"UT{i:03d}"
        sheet = re.sub(r"[\\/?*\[\]:']", "_", f"{fn.cls}.{fn.name}")
        if len(sheet) > 31:
            sheet = re.sub(r"(Service|Controller)\.", ".", sheet)[:31]
        while sheet.lower() in used:
            sheet = f"{sheet[:27]}_{i:03d}"
        used.add(sheet.lower())
        fn.sheet = sheet
    return ordered


def jacoco_totals() -> tuple[float, float] | None:
    if not JACOCO_CSV.exists():
        return None
    rows = list(csv.DictReader(JACOCO_CSV.open(encoding="utf-8")))
    total = lambda k: sum(int(r[k]) for r in rows)  # noqa: E731
    line = total("LINE_COVERED") * 100 / (total("LINE_COVERED") + total("LINE_MISSED"))
    branch = total("BRANCH_COVERED") * 100 / (total("BRANCH_COVERED") + total("BRANCH_MISSED"))
    return round(line, 2), round(branch, 2)


def test_feature_map() -> dict[str, set[str]]:
    if not MATRIX.exists():
        return {}
    from openpyxl import load_workbook
    out = defaultdict(set)
    for row in load_workbook(MATRIX, read_only=True)["4_Existing Tests"].iter_rows(min_row=2, values_only=True):
        if row[0] and row[5] and row[5] != "CHƯA ÁNH XẠ":
            out[Path(row[0]).stem] |= set(row[5].split(", "))
    return out


# ------------------------------------------------------------------ excel --
def serial(when: datetime) -> int:
    return (when.date() - datetime(1899, 12, 30).date()).days


def ensure_rows(ws, at: int, have: int, need: int) -> int:
    """Insert rows before `at` (copying the format of row at-1) so a block has `need` rows; returns added."""
    extra = max(0, need - have)
    if extra:
        ws.Rows(at - 1).Copy()
        ws.Range(ws.Rows(at), ws.Rows(at + extra - 1)).Insert(XL_SHIFT_DOWN)
    return extra


def fit_rows(ws, first: int, have: int, need: int) -> int:
    """Resize the block of `have` rows starting at `first` to `need` rows (min 1); returns the row delta."""
    need = max(need, 1)
    if need < have:
        ws.Range(ws.Rows(first + need), ws.Rows(first + have - 1)).Delete()
        return need - have
    return ensure_rows(ws, first + have, have, need)


def filled_rows(ws, first: int, col: int) -> int:
    r = first
    while ws.Cells(r, col).Value not in (None, ""):
        r += 1
    return r - first


def clear(rng) -> None:
    """ClearContents that tolerates merged cells."""
    if not rng.MergeCells and rng.MergeCells is not None:
        rng.ClearContents()
        return
    for cell in rng.Cells:
        if cell.MergeCells:
            if cell.Address == cell.MergeArea.Cells(1, 1).Address:
                cell.MergeArea.ClearContents()
        else:
            cell.ClearContents()


def write_function(ws, fn: Function, row_in_list: int, overrides: dict) -> None:
    n = len(fn.cases)
    if n > 15:
        ws.Columns(20).Copy()
        ws.Range(ws.Columns(21), ws.Columns(20 + n - 15)).Insert(XL_SHIFT_RIGHT)
    last_col = 5 + max(n, 15)
    for rng in ("D11:D13", "B14:D30", "D32:D33", "D35:D35", "D37:D38"):
        clear(ws.Range(rng))
    clear(ws.Range(ws.Cells(9, 6), ws.Cells(42, last_col)))

    ws.Range("C2,L2").NumberFormat = "General"
    ws.Range("C2").Formula = f"=Functions!E{row_in_list}"
    ws.Range("L2").Formula = f"=Functions!D{row_in_list}"
    authors = Counter(c.author for c in fn.cases).most_common()
    ws.Range("C3").Value = ", ".join(a for a, _ in authors[:2])
    ws.Range("L3").Value = EXECUTED_BY
    ws.Range("C4").Value = fn.loc if fn.loc else ""
    ws.Range("C5").Value = overrides.get("test_requirement") or (
        f"Verify {fn.cls}.{fn.name}: normal, abnormal and boundary behaviour covered by "
        + ", ".join(sorted({c.test_class for c in fn.cases})) + ".")

    for i, case in enumerate(fn.cases):
        ws.Cells(9, 6 + i).Value = f"UTCID{i + 1:02d}"

    # Condition: precondition (row 11), scenarios (rows 15..)
    ws.Range("D11").Value = fn.precondition
    ws.Range(ws.Cells(11, 6), ws.Cells(11, 5 + n)).Value = [["O"] * n]
    ws.Range("B14").Value = "Scenario"
    shift = ensure_rows(ws, 31, 16, n)
    scenario_over = overrides.get("scenarios", {})
    for i, case in enumerate(fn.cases):
        ws.Cells(15 + i, 4).Value = scenario_over.get(case.method, case.scenario)
        ws.Cells(15 + i, 6 + i).Value = "O"

    # Confirm: return values, exceptions
    returns = list(dict.fromkeys(c.returns for c in fn.cases if c.returns))
    exceptions = list(dict.fromkeys(c.exception for c in fn.cases if c.exception))
    ret_row = 32 + shift
    shift += ensure_rows(ws, 34 + shift, 2, len(returns))
    exc_row = 35 + shift
    shift += ensure_rows(ws, 36 + shift, 1, len(exceptions))
    for values, first, attr in ((returns, ret_row, "returns"), (exceptions, exc_row, "exception")):
        for j, value in enumerate(values):
            ws.Cells(first + j, 4).Value = value
            for i, case in enumerate(fn.cases):
                if getattr(case, attr) == value:
                    ws.Cells(first + j, 6 + i).Value = "O"

    # Result rows
    result = 39 + shift
    ws.Range(ws.Cells(result, 6), ws.Cells(result, 5 + n)).Value = [[c.kind for c in fn.cases]]
    ws.Range(ws.Cells(result + 1, 6), ws.Cells(result + 1, 5 + n)).Value = [[c.status for c in fn.cases]]
    dates = ws.Range(ws.Cells(result + 2, 6), ws.Cells(result + 2, 5 + n))
    dates.Value = [[serial(c.executed) if c.executed else "" for c in fn.cases]]
    dates.NumberFormat = "dd/mm/yyyy"
    ws.Range(ws.Cells(15, 4), ws.Cells(exc_row + max(len(exceptions), 1), 4)).WrapText = True
    ws.Columns(4).ColumnWidth = 45


def write_cover(ws, cover: dict, history: list[dict]) -> None:
    latest = history[-1] if history else {}
    ws.Range("B4").Value = cover.get("project_name") or ""
    ws.Range("F4").Value = cover.get("creator") or ""
    ws.Range("B5").Value = cover.get("project_code") or ""
    ws.Range("F5").Value = latest.get("date") or cover.get("issue_date") or ""
    ws.Range("F6").Value = latest.get("version") or cover.get("version") or ""
    ws.Range("B6").Formula = '=B5&"_Unit Test_"&F6'
    fit_rows(ws, 11, max(filled_rows(ws, 11, 2), 1), len(history))
    ws.Range(ws.Cells(11, 1), ws.Cells(10 + max(len(history), 1), 6)).ClearContents()
    for i, change in enumerate(history):
        ws.Range(ws.Cells(11 + i, 1), ws.Cells(11 + i, 6)).Value = [[str(change.get(k) or "") for k in HISTORY_KEYS]]


def read_history(ws, cover: dict) -> list[dict]:
    rows = []
    r = 11
    while ws.Cells(r, 2).Value:
        if not str(ws.Cells(r, 1).Value or "").startswith("<"):
            rows.append({k: ws.Cells(r, c).Text for c, k in enumerate(HISTORY_KEYS, 1)})
        r += 1
    return rows or [{**c, "description": c["description"].replace("system test cases for 53 features",
                                                                  "unit test cases from JUnit tests"),
                     "item": "Unit Test", "reference": "backend/src/test/java, JaCoCo report"} for c in cover.get("changes", [])]


def write_lists(wb, functions: list[Function], cover: dict) -> None:
    ws = wb.Worksheets("Functions")
    ws.Range("E7").Value = ENVIRONMENT
    for h in list(ws.Hyperlinks):
        h.Delete()
    fit_rows(ws, 11, max(filled_rows(ws, 11, 1), 3), len(functions))
    ws.Range(ws.Cells(11, 1), ws.Cells(10 + len(functions), 8)).ClearContents()
    data = [[i, ", ".join(sorted(fn.features)[:3]), fn.cls, fn.name, fn.code, fn.sheet,
             f"{len(fn.cases)} unit test cases from " + ", ".join(sorted({c.test_class for c in fn.cases})),
             fn.precondition] for i, fn in enumerate(functions, 1)]
    ws.Range(ws.Cells(11, 1), ws.Cells(10 + len(functions), 8)).Value = data
    for i, fn in enumerate(functions):
        ws.Hyperlinks.Add(ws.Cells(11 + i, 6), "", f"'{fn.sheet}'!A1", "", fn.sheet)

    ws = wb.Worksheets("Statistics")
    for h in list(ws.Hyperlinks):
        h.Delete()
    ws.Range("B6").Formula = "=Cover!B6"
    ws.Range("F4").Value = cover.get("creator") or ""
    ws.Range("F5").Value = cover.get("reviewer") or ""
    ws.Range("F6").Formula = "=Cover!F5"
    classes = len({fn.cls for fn in functions})
    total = sum(len(fn.cases) for fn in functions)
    ws.Range("B7").Value = (f"This release covers {len(functions)} functions in {classes} classes "
                            f"with {total} unit test cases (JUnit 5).")
    sub_total = ws.Columns(2).Find("Sub total").Row
    fit_rows(ws, 12, sub_total - 14, len(functions))
    ws.Range(ws.Cells(12, 1), ws.Cells(11 + len(functions), 9)).ClearContents()
    rows = []
    for i, fn in enumerate(functions, 1):
        ref = f"'{fn.sheet}'!"
        rows.append([i, fn.code] + [f"={ref}{c}7" for c in ("A", "C", "F", "L", "M", "N", "O")])
    ws.Range(ws.Cells(12, 1), ws.Cells(11 + len(functions), 9)).Formula = rows
    for i, fn in enumerate(functions):
        ws.Hyperlinks.Add(ws.Cells(12 + i, 2), "", f"'{fn.sheet}'!A1", "", fn.code)
    boundary_row = ws.Columns(2).Find("Boundary case").Row
    totals = jacoco_totals()
    if totals:
        for j, (label, value) in enumerate((("Line coverage (JaCoCo)", totals[0]),
                                            ("Branch coverage (JaCoCo)", totals[1])), 1):
            ws.Rows(boundary_row).Copy()
            ws.Rows(boundary_row + j).PasteSpecial(XL_PASTE_FORMATS)
            ws.Cells(boundary_row + j, 2).Value = label
            ws.Cells(boundary_row + j, 4).Value = value
            ws.Cells(boundary_row + j, 5).Value = "%"


def snapshot_of(wb) -> dict[str, str]:
    names = [ws.Name for ws in wb.Worksheets]
    if SNAPSHOT_SHEET not in names:
        return {}
    ws = wb.Worksheets(SNAPSHOT_SHEET)
    values = ws.UsedRange.Value
    return {str(r[0]): str(r[1] or "") for r in values[1:] if r and r[0]} if values else {}


def write_snapshot(wb, snapshot: dict[str, str]) -> None:
    names = [ws.Name for ws in wb.Worksheets]
    ws = wb.Worksheets(SNAPSHOT_SHEET) if SNAPSHOT_SHEET in names else wb.Worksheets.Add(
        None, wb.Worksheets(wb.Worksheets.Count))
    ws.Name = SNAPSHOT_SHEET
    ws.Cells.ClearContents()
    rows = [["Test", "Result"]] + [[k, v] for k, v in sorted(snapshot.items())]
    ws.Range(ws.Cells(1, 1), ws.Cells(len(rows), 2)).Value = rows
    ws.Visible = 0


def copy_after(wb, sheet, anchor):
    # pywin32 late binding silently ignores Copy(After=...); pass Before/After positionally.
    count = wb.Sheets.Count
    sheet.Copy(None, anchor)
    if wb.Sheets.Count != count + 1:
        raise RuntimeError(f"Excel did not copy sheet {sheet.Name!r}")
    return wb.Sheets(anchor.Index + 1)


# ------------------------------------------------------------------- main --
def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true", help="analyse tests only, do not touch the workbook")
    parser.add_argument("--note", default="", help="extra text for the new Record of change entry")
    parser.add_argument("--report", type=Path, default=REPORT)
    args = parser.parse_args()

    classes = {p.stem: p for p in MAIN.rglob("*.java")}
    functions = analyse(classes, test_feature_map())
    cases = [c for fn in functions for c in fn.cases]
    kinds = Counter(c.kind for c in cases)
    states = Counter(c.status or "Untested" for c in cases)
    print(f"Functions: {len(functions)} | UTCIDs: {len(cases)} | N/A/B: {kinds['N']}/{kinds['A']}/{kinds['B']} | "
          f"P/F/Untested: {states['P']}/{states['F']}/{states['Untested']} | "
          f"max UTCID per function: {max(len(f.cases) for f in functions)}")
    if args.dry_run:
        for fn in functions:
            print(f"  {fn.code} {fn.sheet:32} {len(fn.cases):3} cases  LOC={fn.loc}")
        return

    report = args.report.resolve()
    try:
        with open(report, "r+b"):
            pass
    except PermissionError:
        raise SystemExit(f"{report} is open in another program (Excel?). Close it and run again.")
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(report, BACKUP_DIR / f"Report5_Unit Test_{datetime.now():%Y%m%d_%H%M%S}.xls")

    import win32com.client
    cover = load_content("_cover")
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible, xl.DisplayAlerts, xl.ScreenUpdating = False, False, False
    wb = xl.Workbooks.Open(str(report))
    try:
        xl.Calculation = -4135  # manual while writing
        names = [ws.Name for ws in wb.Worksheets]
        if TEMPLATE_SHEET not in names:
            copy_after(wb, wb.Worksheets("Function 1"), wb.Worksheets(wb.Worksheets.Count)).Name = TEMPLATE_SHEET
        tpl = wb.Worksheets(TEMPLATE_SHEET)
        tpl.Visible = 0
        history = read_history(wb.Worksheets("Cover"), cover)
        before = snapshot_of(wb)

        keep = set(FIXED_SHEETS) | {TEMPLATE_SHEET, SNAPSHOT_SHEET}
        for ws in list(wb.Worksheets):
            if ws.Name not in keep:
                ws.Delete()

        anchor = wb.Worksheets("Statistics")
        for i, fn in enumerate(functions):
            tpl.Visible = -1
            ws = copy_after(wb, tpl, anchor)
            tpl.Visible = 0
            ws.Name = fn.sheet
            write_function(ws, fn, 11 + i, load_content_unit(fn.cls).get(fn.name, {}))
            anchor = ws

        write_lists(wb, functions, cover)
        after = {f"{c.test_class}.{c.method}": {"P": "Passed", "F": "Failed"}.get(c.status, "Untested")
                 for c in cases}
        new_version = next_version(history, before, after, args.note, labels=("Passed", "Failed", "Untested"),
                                   item="Unit test cases", result_name="results",
                                   reference="backend/src/test/java, JaCoCo report") if before else None
        if new_version:
            history.append(new_version)
        write_cover(wb.Worksheets("Cover"), cover, history)
        write_snapshot(wb, after)
        wb.Worksheets("Cover").Activate()
        xl.Calculation = -4105
        wb.Save()
    finally:
        wb.Close(False)
        xl.Quit()
    print(f"Version: {history[-1]['version']} ({'new: ' + history[-1]['description'] if new_version else 'unchanged'})")
    print(f"Saved: {report}")


def load_content_unit(cls: str) -> dict:
    path = CONTENT_DIR / f"{cls}.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


if __name__ == "__main__":
    main()
