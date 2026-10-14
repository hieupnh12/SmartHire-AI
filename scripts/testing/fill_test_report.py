"""Fill docs/test/Report5_Test Report.xlsx from the traceability matrix.

Re-runnable: rows are upserted by Test Case ID. Existing values in the result columns
(Round 1..3, Test date, Tester, Note) are kept unless --refresh-results is given.
Business-rule test cases are written only when their content exists in
scripts/testing/content/test_report/<CODE>.json; the hidden "_Progress" sheet shows what is left.

Usage:
  python scripts/testing/build_traceability.py
  python scripts/testing/fill_test_report.py [--refresh-results]
"""
from __future__ import annotations

import argparse
import json
import math
import re
import shutil
import subprocess
import xml.etree.ElementTree as ET
from collections import defaultdict
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "docs" / "test" / "Report5_Test Report.xlsx"
MATRIX = ROOT / ".tmp" / "testing" / "Traceability_Matrix.xlsx"
CONTENT_DIR = Path(__file__).resolve().parent / "content" / "test_report"
SUREFIRE = ROOT / "backend" / "target" / "surefire-reports"
FEATURE_README = ROOT / "docs" / "features" / "README.md"
BACKUP_DIR = ROOT / ".tmp" / "testing" / "backup"

FIRST_ROW = 11
RESULT_COLS = range(6, 16)  # F..O
TEMPLATE_SHEET = "_Template"
PROGRESS_SHEET = "_Progress"

# Ordered: the first matching prefix wins.
OWNERS = [
    ("MASTER-08", "Nguyễn Minh Hiếu"),
    ("MASTER-", "Nguyễn Nhất Sinh"),
    ("AUTH-", "Nguyễn Nhất Sinh"),
    ("COMPANY-", "Lê Công Cường"),
    ("CV-01", "Lê Phúc Tiến"),
    ("CV-06", "Nguyễn Minh Hiếu"),
    ("JOB-", "Lê Phúc Tiến"),
    ("CV-", "Lê Công Cường"),
    ("SCHED-02", "Lê Công Cường"),
    ("SCHED-03", "Lê Công Cường"),
    ("SCHED-01", "Lê Mỹ Lộc"),
    ("ASSESS-", "Lê Mỹ Lộc"),
    ("INT-", "Lê Mỹ Lộc"),
    ("PRACT-", "Lê Mỹ Lộc"),
    ("RANK-", "Nguyễn Minh Hiếu"),
    ("WF-", "Nguyễn Minh Hiếu"),
    ("DASH-", "Nguyễn Minh Hiếu"),
]
NAME_OVERRIDES = {
    "COMPANY-02-Tenant-Landing-Page": "Tenant Landing / Career Page",
    "MASTER-04": "Self-serve Checkout & Click-wrap Legal Consent (Decree 13/2023/ND-CP)",
}

# Mirrors the permitAll rules in SecurityConfig.
PUBLIC_RULES = [
    (None, re.compile(r"^/api/v1/public/")),
    ("POST", re.compile(r"^/api/v1/(master|tenant)/auth/(login|refresh|logout|forgot-password|reset-password|google)$")),
    ("POST", re.compile(r"^/api/v1/master/consultations$")),
    ("GET", re.compile(r"^/api/v1/master/tenants/check(-subdomain)?/\{\}$")),
    ("POST", re.compile(r"^/api/v1/tenant/users/invitations/accept$")),
]

ENVIRONMENT = (
    "1. Server: Spring Boot 3 backend (Java 21) and React 19 frontend served by Nginx (Docker Compose).\n"
    "2. Database: PostgreSQL master DB (smarthire_master) and one MySQL database per tenant; Redis; RabbitMQ.\n"
    "3. Web browser: latest Google Chrome / Microsoft Edge.\n"
    "4. Test data: at least two ACTIVE tenants (tenant A, tenant B) with Tenant Admin, HR, Recruiter and "
    "Candidate accounts; one Workspace Admin account on the master platform."
)


def owner_of(code: str) -> str:
    return next((name for prefix, name in OWNERS if code.startswith(prefix)), "")


def is_public(method: str, path: str) -> bool:
    return any((m is None or m == method) and rx.search(path) for m, rx in PUBLIC_RULES)


def humanize(handler: str) -> str:
    words = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", handler).split()
    return " ".join(words).capitalize() if words else handler


# ------------------------------------------------------------------ inputs --
def read_sheet(wb, name: str) -> list[tuple]:
    return [r for r in wb[name].iter_rows(min_row=2, values_only=True) if r[0]]


def english_names() -> dict[str, str]:
    names = {}
    for m in re.finditer(r"^\|\s*([A-Z]+-\d+)\s*\|\s*([^|]+?)\s*\|", FEATURE_README.read_text(encoding="utf-8"), re.M):
        names[m.group(1)] = m.group(2)
    return names


ROLE_WORDS = ["Workspace Admin", "Tenant Admin", "Admin", "HR", "Recruiter", "Hiring Manager",
              "Interviewer", "Staff", "Candidate"]


def doc_actor(code: str, doc: str) -> str:
    """English role names found in the doc's Actor section (docs mix Vietnamese prose into it)."""
    if code.startswith("MASTER"):
        return "Workspace Admin"
    text = (ROOT / doc).read_text(encoding="utf-8")
    m = re.search(r"^#+\s*Actor\s*\n+((?:\s*[-*].*\n?)+)", text, re.M)
    section = m.group(1) if m else ""
    found = []
    for role in ROLE_WORDS:
        if re.search(r"\b" + role + r"\b", section, re.I) and not any(role in f for f in found):
            found.append(role)
    return found[0] if found else "Recruiter"


def load_surefire() -> dict[str, tuple[list[str], datetime]]:
    """test class simple name -> (failed test names, run time)."""
    results = {}
    for path in SUREFIRE.glob("TEST-*.xml"):
        root = ET.parse(path).getroot()
        failed = [f"{root.get('name').split('.')[-1]}.{tc.get('name')}" for tc in root.iter("testcase")
                  if tc.find("failure") is not None or tc.find("error") is not None]
        stamp = root.get("timestamp")
        when = datetime.fromisoformat(stamp) if stamp else datetime.fromtimestamp(path.stat().st_mtime)
        results[root.get("name").split(".")[-1]] = (failed, when)
    return results


def test_commit_dates() -> dict[str, datetime]:
    """test class simple name -> date of the last commit touching it (CI runs mvn test on every push)."""
    log = subprocess.run(
        ["git", "log", "--name-only", "--format=@%ad", "--date=short", "--", "backend/src/test/java"],
        cwd=ROOT, capture_output=True, text=True, encoding="utf-8", check=True).stdout
    dates: dict[str, datetime] = {}
    current = None
    for line in log.splitlines():
        if line.startswith("@"):
            current = datetime.strptime(line[1:], "%Y-%m-%d")
        elif line.endswith(".java"):
            dates.setdefault(Path(line).stem, current)
    return dates


def load_content(code: str) -> dict:
    path = CONTENT_DIR / f"{code}.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


# --------------------------------------------------------------- test cases --
def api_cases(code: str, actor: str, api_rows: list[tuple]) -> list[tuple[str, dict]]:
    """Returns (function group, case) pairs for API and security test cases."""
    out = []
    protected, examples = [], []
    for _, method, path, impl, tc_ok, tc_val, tc_auth, tc_tenant in api_rows:
        if method == "*":
            continue
        public = is_public(method, path)
        path = path.replace("{}", "{id}")
        if not public:
            protected.append(f"{method} {path}")
            if "{id}" in path:
                examples.append(f"{method} {path}")
        handler = impl.split(".")[-1].split(",")[0] if impl != "KHÔNG TÌM THẤY" else ""
        title = humanize(handler) if handler else f"{method} {path}"
        group = f"{title} ({method} {path})"
        master = code.startswith("MASTER") or path.startswith("/api/v1/master")
        login = ("1. Do not log in (public endpoint)." if public
                 else f"1. Log in as {actor}." if master
                 else f"1. Log in as {actor} of tenant A.")
        header = "" if public or master else " with header X-Tenant-ID = tenant A"
        body = " and a valid request body" if method in ("POST", "PUT", "PATCH") else ""
        expected = {
            "GET": ("HTTP 200. Response success = true and contains the requested data."
                    if master or public else
                    "HTTP 200. Response success = true and contains the requested data of tenant A only."),
            "POST": "HTTP 200/201. Response success = true; the record/action is saved and returned in the response.",
            "PUT": "HTTP 200. Response success = true; the changes are saved and returned.",
            "PATCH": "HTTP 200. Response success = true; the changed fields are saved.",
            "DELETE": "HTTP 200/204. The record is removed (or archived) and no longer appears in lists.",
        }[method]
        pre = "Master platform is running." if master else "Tenant A is ACTIVE."
        if not public:
            pre += f" The {actor} account has permission for this action."
        if "{id}" in path:
            pre += " The target resource exists."
        out.append((group, {
            "id": tc_ok, "impl": impl,
            "description": f"Verify {title.lower()} succeeds with valid input ({method} {path}).",
            "procedure": f"{login}\n2. Send {method} {path}{header}{body}.\n3. Check the response and the data shown in the system.",
            "expected": expected, "precondition": pre,
        }))
        if tc_val:
            out.append((group, {
                "id": tc_val, "impl": impl,
                "description": f"Verify {title.lower()} rejects invalid input ({method} {path}).",
                "procedure": (f"{login}\n2. Send {method} {path}{header} with missing required fields or invalid "
                              "values (empty text, wrong format, out-of-range numbers, non-existing IDs).\n"
                              "3. Check the response and the data in the system."),
                "expected": ("HTTP 400 (or 404 for a non-existing ID) with success = false and an error message / "
                             "field errors. No data is created or changed."),
                "precondition": pre,
            }))

    security = [r for r in api_rows if r[1] == "*"]
    if security:
        _, _, _, _, _, _, tc_auth, tc_tenant = security[0]
        group = "Security & tenant isolation"
        target = protected[0] if protected else ""
        if code.startswith("MASTER"):
            auth_case = {
                "description": "Verify platform admin APIs reject unauthenticated requests and tenant users.",
                "procedure": (f"1. Send {target} without an Authorization header.\n"
                              "2. Send the same request with a valid token of a tenant user (e.g. Tenant Admin)."),
                "expected": "Step 1: HTTP 401, code UNAUTHORIZED.\nStep 2: HTTP 403, code FORBIDDEN (requires WORKSPACE_ADMIN).",
                "precondition": "A tenant user account exists.",
            }
            tenant_case = {
                "description": "Verify a Workspace Admin token cannot be used to read tenant data through tenant APIs.",
                "procedure": ("1. Log in as Workspace Admin.\n2. Call a tenant API (e.g. GET /api/v1/jobs) with header "
                              "X-Tenant-ID = tenant A using the master token."),
                "expected": "The request is rejected (HTTP 401/403) and no tenant data is returned.",
                "precondition": "Workspace Admin account exists; tenant A has data.",
            }
        else:
            auth_case = {
                "description": "Verify protected APIs of this feature reject unauthenticated and unauthorized users.",
                "procedure": (f"1. Send {target} without an Authorization header.\n"
                              "2. Send the same request with a valid token of a role that is not allowed "
                              "(e.g. Candidate calling a staff API)."),
                "expected": "Step 1: HTTP 401, code UNAUTHORIZED.\nStep 2: HTTP 403, code FORBIDDEN. No data is returned or changed.",
                "precondition": "Accounts of an allowed role and a not-allowed role exist in tenant A.",
            }
            example = examples[0] if examples else target
            tenant_case = {
                "description": "Verify tenant data isolation: a tenant B user cannot read or change tenant A data.",
                "procedure": (f"1. Log in as {actor} of tenant B.\n2. Send {example} using a resource ID that "
                              "belongs to tenant A, with header X-Tenant-ID = tenant A.\n"
                              "3. Open the related list screen of tenant B."),
                "expected": ("Step 2: the tenant in the JWT (tenant B) is used, so the request returns HTTP 404 "
                             "(or an empty result); no tenant A data is returned or changed.\n"
                             "Step 3: only tenant B records are listed."),
                "precondition": "Tenant A and tenant B are ACTIVE and both have data for this feature.",
            }
        if protected:
            out.append((group, {"id": tc_auth, "impl": "SecurityConfig", **auth_case}))
            out.append((group, {"id": tc_tenant, "impl": "JwtAuthenticationFilter", **tenant_case}))
        else:
            na = {"impl": "public", "na": "All endpoints of this feature are public (permitAll)."}
            out.append((group, {"id": tc_auth, **na, **auth_case}))
            out.append((group, {"id": tc_tenant, **na, **tenant_case}))
    return out


def rule_cases(code: str, rules: list[tuple], content: dict) -> tuple[list[tuple[str, dict]], list[str]]:
    written, warnings = [], []
    cases = content.get("cases", {})
    for _, br_id, rule, tc_id, _, _ in rules:
        item = cases.get(tc_id)
        if not item:
            continue
        if item.get("rule") and item["rule"].strip() != rule.strip():
            warnings.append(f"{tc_id}: business rule text changed in docs, review content ({br_id})")
        written.append((item.get("function", "Business rules"), {"id": tc_id, "impl": "", **item}))
    return written, warnings


# ------------------------------------------------------------------ writer --
def black(cell) -> None:
    font = copy(cell.font)
    font.color = "FF000000"
    cell.font = font


def copy_row_style(ws, src_ws, src_row: int, dst_row: int, max_col: int = 15) -> None:
    for col in range(1, max_col + 1):
        src = src_ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        black(dst)


def row_height(ws, values: list) -> float:
    lines = 1
    for col, value in enumerate(values, 1):
        if not value:
            continue
        width = ws.column_dimensions[chr(64 + col)].width or 10
        chars = max(int(width * 1.15), 8)
        lines = max(lines, sum(max(1, math.ceil(len(part) / chars)) for part in str(value).split("\n")))
    return min(409, 15 * lines + 4)


def prepare_template(wb) -> None:
    if TEMPLATE_SHEET not in wb.sheetnames:
        tpl = wb.copy_worksheet(wb["Feature 1"])
        tpl.title = TEMPLATE_SHEET
        tpl.sheet_state = "hidden"
    for name in ("Feature 1", "Feature 2"):
        if name in wb.sheetnames:
            del wb[name]


def write_feature(wb, code: str, name: str, requirement: str, groups: list[tuple[str, dict]],
                  results: dict[str, tuple[str, str, str, str]], refresh: bool) -> None:
    tpl = wb[TEMPLATE_SHEET]
    if code in wb.sheetnames:
        ws = wb[code]
    else:
        ws = wb.copy_worksheet(tpl)
        ws.title = code
        ws.sheet_state = "visible"

    existing = {}
    for row in ws.iter_rows(min_row=FIRST_ROW + 1, values_only=True):
        if row[0] and "-TC" in str(row[0]):
            existing[str(row[0])] = row[5:15]
    ws.delete_rows(FIRST_ROW, max(ws.max_row - FIRST_ROW + 1, 1))

    ws["B2"] = name
    ws["B3"] = requirement
    black(ws["B2"])
    black(ws["B3"])
    ws["B4"] = f'=COUNTIF(A{FIRST_ROW}:A2000,"*-TC*")'
    for row, col in ((6, "F"), (7, "I"), (8, "L")):
        for c in "BCDE":
            ws[f"{c}{row}"] = f"=COUNTIF(${col}${FIRST_ROW}:${col}$2000,{c}$5)"

    r = FIRST_ROW
    current = None
    for group, case in groups:
        if group != current:
            copy_row_style(ws, tpl, 11, r)
            ws.cell(r, 1, group)
            ws.row_dimensions[r].height = 15.75
            current, r = group, r + 1
        copy_row_style(ws, tpl, 12, r)
        values = [case["id"], case["description"], case["procedure"], case["expected"], case["precondition"]]
        for col, value in enumerate(values, 1):
            ws.cell(r, col, value)
        kept = existing.get(case["id"])
        if kept and not refresh:
            for col, value in zip(RESULT_COLS, kept):
                ws.cell(r, col, value)
        else:
            status, date, tester, note = results[case["id"]]
            for col, value in zip((6, 7, 8, 15), (status, date, tester, note)):
                ws.cell(r, col, value)
        ws.row_dimensions[r].height = row_height(ws, values)
        r += 1

    ws.data_validations.dataValidation = []
    dv = DataValidation(type="list", formula1="$R$2:$R$5", allow_blank=True)
    for col in "FIL":
        dv.add(f"{col}{FIRST_ROW}:{col}2000")
    ws.add_data_validation(dv)


HISTORY_KEYS = ("date", "version", "item", "type", "description", "reference")


def read_history(wb, cover: dict) -> list[dict]:
    """Record of change rows already in the Cover sheet; falls back to _cover.json on first fill."""
    ws = wb["Cover"]
    rows = []
    for row in ws.iter_rows(min_row=11, max_col=6, values_only=True):
        if row[1] and not str(row[0] or "").startswith("<"):
            rows.append(dict(zip(HISTORY_KEYS, row)))
    return rows or list(cover.get("changes", []))


def results_snapshot(wb) -> dict[str, str]:
    """Test Case ID -> Round 1 result for every feature sheet."""
    snapshot = {}
    for ws in wb.worksheets:
        if ws.title.startswith("_") or ws.title in ("Cover", "Test Cases", "Test Statistics"):
            continue
        for row in ws.iter_rows(min_row=FIRST_ROW + 1, max_col=6, values_only=True):
            if row[0] and "-TC" in str(row[0]):
                snapshot[str(row[0])] = row[5] or ""
    return snapshot


def next_version(history: list[dict], before: dict[str, str], after: dict[str, str], note: str,
                 labels: tuple[str, ...] = ("Passed", "Failed", "Pending", "N/A"),
                 item: str = "Test cases", result_name: str = "Round 1 results",
                 reference: str = "SmartHire-AI feature specifications (docs/features)") -> dict | None:
    added = after.keys() - before.keys()
    removed = before.keys() - after.keys()
    changed = [k for k in after.keys() & before.keys() if after[k] != before[k]]
    if not (added or removed or changed or note):
        return None
    majors = [int(m.group(1)) for h in history if (m := re.match(r"v?(\d+)", str(h.get("version") or "")))]
    parts = []
    if added:
        parts.append(f"Added {len(added)} test cases")
    if removed:
        parts.append(f"Removed {len(removed)} test cases")
    if changed:
        parts.append(f"Updated {result_name} of {len(changed)} test cases")
    totals = {s: sum(1 for v in after.values() if v == s) for s in labels}
    parts.append(f"Total {len(after)} test cases (" + ", ".join(f"{k} {v}" for k, v in totals.items()) + ")")
    if note:
        parts.append(note)
    return {
        "date": f"{datetime.now():%d/%m/%Y}",
        "version": f"v{(max(majors) if majors else 0) + 1}.0",
        "item": item,
        "type": ", ".join(t for t, flag in (("A", added), ("M", changed or note), ("D", removed)) if flag),
        "description": "; ".join(parts) + ".",
        "reference": reference,
    }


def write_cover(wb, cover: dict, history: list[dict]) -> None:
    ws = wb["Cover"]
    latest = history[-1] if history else {}
    for ref, value in (("B4", cover.get("project_name")), ("F4", cover.get("creator")),
                       ("B5", cover.get("project_code")),
                       ("F5", latest.get("date") or cover.get("issue_date")),
                       ("F6", latest.get("version") or cover.get("version"))):
        ws[ref] = value or None
        black(ws[ref])
    ws["B6"] = '=B5&"_Test Report_"&F6'
    black(ws["B6"])
    style = [copy(ws.cell(11, c)._style) for c in range(1, 7)]
    for row in ws.iter_rows(min_row=11, max_row=max(ws.max_row, 17)):
        for cell in row:
            cell.value = None
    for i, change in enumerate(history):
        for col, key in enumerate(HISTORY_KEYS, 1):
            cell = ws.cell(11 + i, col, change.get(key))
            cell._style = copy(style[col - 1])
            black(cell)


def write_index_sheets(wb, features: list[tuple[str, str, str, str]], cover: dict) -> None:
    ws = wb["Test Cases"]
    ws["D3"], ws["D4"], ws["D5"] = "=Cover!B4", '=IF(Cover!B5="","",Cover!B5)', ENVIRONMENT
    for ref in ("D3", "D4", "D5"):
        black(ws[ref])
    styles = [copy(ws.cell(9, c)._style) for c in range(2, 7)]
    for row in ws.iter_rows(min_row=9, max_row=max(ws.max_row, 9)):
        for cell in row:
            cell.value = None
    for i, (code, name, requirement, precondition) in enumerate(features, 1):
        values = [i, name, code, requirement, precondition]
        for col, value in enumerate(values, 2):
            cell = ws.cell(8 + i, col, value)
            cell._style = copy(styles[col - 2])
            black(cell)

    ws = wb["Test Statistics"]
    row_style = [copy(ws.cell(11, c)._style) for c in range(2, 9)]
    total_style = [copy(ws.cell(14, c)._style) for c in range(2, 9)]
    cov_style = [copy(ws.cell(16, c)._style) for c in range(2, 9)]
    ws["C3"], ws["C4"] = "=Cover!B4", '=IF(Cover!B5="","",Cover!B5)'
    ws["C5"] = '=Cover!B6'
    ws["H3"], ws["H4"], ws["H5"] = cover.get("creator") or None, cover.get("reviewer") or None, "=Cover!F5"
    ws["C6"] = (f"This release includes {len(features)} features: "
                f"{sum(1 for f in features if f[0].startswith('MASTER'))} Master platform modules and "
                f"{sum(1 for f in features if not f[0].startswith('MASTER'))} Tenant modules.")
    for ref in ("C3", "C4", "C5", "C6", "H3", "H4", "H5"):
        black(ws[ref])
    for row in ws.iter_rows(min_row=11, max_row=max(ws.max_row, 20)):
        for cell in row:
            cell.value = None
    first, last = 11, 10 + len(features)
    for i, (code, *_) in enumerate(features):
        r = first + i
        values = [i + 1, f"='{code}'!B2", f"='{code}'!B6", f"='{code}'!C6", f"='{code}'!D6",
                  f"='{code}'!E6", f"='{code}'!B4"]
        for col, value in enumerate(values, 2):
            ws.cell(r, col, value)._style = copy(row_style[col - 2])
            black(ws.cell(r, col))
    total = last + 2
    ws.cell(total, 3, "Sub total")
    for col, letter in zip(range(4, 9), "DEFGH"):
        ws.cell(total, col, f"=SUM({letter}{first}:{letter}{last})")
    for col in range(2, 9):
        ws.cell(total, col)._style = copy(total_style[col - 2])
        black(ws.cell(total, col))
    for offset, (label, formula) in enumerate((
            ("Test coverage", f"=(D{total}+E{total})*100/(H{total}-G{total})"),
            ("Test successful coverage", f"=D{total}*100/(H{total}-G{total})"))):
        r = total + 2 + offset
        for col in range(2, 9):
            ws.cell(r, col)._style = copy(cov_style[col - 2])
            black(ws.cell(r, col))
        ws.cell(r, 3, label)
        ws.cell(r, 5, formula)
        ws.cell(r, 6, "%")


def write_progress(wb, rows: list[list]) -> None:
    if PROGRESS_SHEET in wb.sheetnames:
        del wb[PROGRESS_SHEET]
    ws = wb.create_sheet(PROGRESS_SHEET)
    ws.sheet_state = "hidden"
    ws.append(["Code", "Owner", "BR planned", "BR written", "API/security TCs", "TCs written",
               "TCs planned", "Round 1", "Filled at"])
    for row in rows:
        ws.append(row)


# -------------------------------------------------------------------- main --
def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--refresh-results", action="store_true",
                        help="recompute Round 1 / date / tester / note even for rows that already exist")
    parser.add_argument("--note", default="", help="extra text for the new Record of change entry")
    parser.add_argument("--report", type=Path, default=REPORT, help="workbook to fill (default: docs/test)")
    args = parser.parse_args()
    refresh = args.refresh_results
    report = args.report

    if not MATRIX.exists():
        raise SystemExit(f"Missing {MATRIX}. Run scripts/testing/build_traceability.py first.")
    matrix = load_workbook(MATRIX, read_only=True)
    summary = read_sheet(matrix, "1_Tong quan")
    rules_by = defaultdict(list)
    for row in read_sheet(matrix, "2_Business Rules"):
        rules_by[row[0]].append(row)
    apis_by = defaultdict(list)
    for row in read_sheet(matrix, "3_API Matrix"):
        apis_by[row[0]].append(row)
    tests_by = defaultdict(set)
    for row in read_sheet(matrix, "4_Existing Tests"):
        for code in (row[5] or "").split(", "):
            tests_by[code].add(Path(row[0]).stem)

    surefire = load_surefire()
    commit_dates = test_commit_dates()
    names = english_names()
    cover = load_content("_cover")
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    try:
        with open(report, "r+b"):
            pass
    except PermissionError:
        raise SystemExit(f"{report} is open in another program (Excel?). Close it and run again.")
    shutil.copy2(report, BACKUP_DIR / f"Report5_Test Report_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    wb = load_workbook(report)
    prepare_template(wb)
    history = read_history(wb, cover)
    before = results_snapshot(wb)

    index, progress, warnings = [], [], []
    now = f"{datetime.now():%d/%m/%Y %H:%M}"
    for code, vi_name, _, _, n_rules, _, _, _, planned, doc in summary:
        content = load_content(code)
        name = content.get("feature_name") or NAME_OVERRIDES.get(code) or names.get(code) or vi_name
        actor = content.get("actor") or doc_actor(code, doc)
        requirement = content.get("test_requirement") or (
            f"Verify the business rules and APIs of {code} - {name} as specified in {doc}.")
        owner = owner_of(code)

        classes = tests_by.get(code, set()) & surefire.keys()
        failed = [f for c in sorted(classes) for f in surefire[c][0]]
        if not classes:
            feature_result = ("Pending", "", owner, "Not tested: no automated test covers this feature yet.")
        elif failed:
            run_at = max(surefire[c][1] for c in classes).strftime("%d/%m/%Y")
            shown = ", ".join(failed[:3]) + (" ..." if len(failed) > 3 else "")
            feature_result = ("Failed", run_at, owner, f"Failing automated tests: {shown}")
        else:
            run_at = max(commit_dates.get(c) or surefire[c][1] for c in classes).strftime("%d/%m/%Y")
            shown = ", ".join(sorted(classes)[:3]) + (" ..." if len(classes) > 3 else "")
            feature_result = ("Passed", run_at, owner, f"Verified by automated tests: {shown}")

        br, warn = rule_cases(code, rules_by.get(code, []), content)
        warnings += warn
        groups = br + api_cases(code, actor, apis_by.get(code, []))
        results = {}
        for _, case in groups:
            if case.get("na"):
                results[case["id"]] = ("N/A", "", owner, case["na"])
            elif case.get("impl") == "KHÔNG TÌM THẤY":
                results[case["id"]] = ("Pending", "", owner, "Not tested: API not implemented yet.")
            else:
                results[case["id"]] = feature_result
        write_feature(wb, code, name, requirement, groups, results, refresh)
        precondition = ("Master platform is running; a Workspace Admin account exists."
                        if code.startswith("MASTER") else
                        f"Tenant A and tenant B are ACTIVE; a {actor} account exists in tenant A.")
        index.append((code, name, requirement, precondition))
        progress.append([code, owner, n_rules, len(br), len(groups) - len(br), len(groups), planned,
                         feature_result[0], now])

    new_version = next_version(history, before, results_snapshot(wb), args.note) if before else None
    if new_version:
        history.append(new_version)
    write_cover(wb, cover, history)
    write_index_sheets(wb, index, cover)
    write_progress(wb, progress)
    order = ["Cover", "Test Cases", "Test Statistics"] + [f[0] for f in index] + [TEMPLATE_SHEET, PROGRESS_SHEET]
    vietnamese = re.compile(r"[\u00C0-\u1EF9\u0110\u0111]")
    tester_cols = {8, 11, 14}
    for sheet in order[:-2]:
        for row in wb[sheet].iter_rows():
            for cell in row:
                if (isinstance(cell.value, str) and vietnamese.search(cell.value)
                        and not (sheet not in ("Cover", "Test Statistics") and cell.column in tester_cols)
                        and cell.value not in (cover.get("creator"), cover.get("reviewer"))):
                    warnings.append(f"Non-English text in {sheet}!{cell.coordinate}: {cell.value[:60]}")
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.save(report)

    written = sum(p[5] for p in progress)
    planned = sum(p[6] for p in progress)
    br_left = [f"{p[0]} ({p[2] - p[3]})" for p in progress if p[3] < p[2]]
    print(f"Sheets: {len(index)} | TCs written: {written}/{planned} | "
          f"Round 1: " + ", ".join(f"{k}={sum(1 for p in progress if p[7] == k)} features"
                                    for k in ("Passed", "Failed", "Pending")))
    print(f"Business-rule content still missing for {len(br_left)} features: {', '.join(br_left)}")
    print(f"Version: {history[-1]['version']} ({'new: ' + history[-1]['description'] if new_version else 'unchanged'})")
    for w in warnings:
        print("WARN", w)
    print(f"Saved: {report}")


if __name__ == "__main__":
    main()
