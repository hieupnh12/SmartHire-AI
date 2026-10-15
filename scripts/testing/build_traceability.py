"""Build a requirement traceability matrix for Report5 (Test Report + Unit Test).

Sources:
  - docs/features/**/*.md          -> features, business rules, documented APIs
  - backend/src/main/java/**       -> controller endpoints, services, public methods
  - backend/src/test/java/**       -> existing @Test methods
  - backend/target/site/jacoco/jacoco.csv (optional) -> per-class coverage

Usage:
  python scripts/testing/build_traceability.py [--out .tmp/testing/Traceability_Matrix.xlsx]
"""
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs" / "features"
MAIN = ROOT / "backend" / "src" / "main" / "java"
TEST = ROOT / "backend" / "src" / "test" / "java"
JACOCO_CSV = ROOT / "backend" / "target" / "site" / "jacoco" / "jacoco.csv"

HTTP = {"GET", "POST", "PUT", "PATCH", "DELETE"}
SUBJECT_SUFFIXES = (
    "Service", "Controller", "Calculator", "Client", "Validator", "Engine", "Scheduler",
    "Generator", "Parser", "Renderer", "Sender", "Filter", "Provider", "Factory", "Listener",
    "Consumer", "Publisher", "Resolver", "Interceptor", "Mapper",
)
# Master docs drifted from code: test cases for these features follow the real controllers.
MASTER_CONTROLLERS = {
    "MASTER-01": {"MasterTenantController"},
    "MASTER-02": {"MasterSubscriptionController"},
    "MASTER-03": {"MasterBillingController", "SePayController", "VnPayController", "PayPalController"},
    "MASTER-04": {"PublicCheckoutController"},
    "MASTER-05": {"MasterContractController", "PublicContractController"},
    "MASTER-06": {"MasterConsultationController"},
    "MASTER-07": {"MasterAiConfigController"},
    "MASTER-08": {"MasterAnalyticsController"},
}
ABNORMAL_WORDS = ("reject", "invalid", "fail", "throw", "forbidden", "denied", "deny", "missing",
                  "notfound", "not_found", "error", "expired", "exceed", "unauthor", "conflict",
                  "duplicate", "block", "locked", "wrong", "without", "cannot", "notallowed", "refuse")
BOUNDARY_WORDS = ("boundary", "limit", "max", "min", "zero", "empty", "exact", "edge", "null",
                  "blank", "threshold", "overflow", "first", "last")


# ---------------------------------------------------------------- features --
@dataclass
class Feature:
    code: str
    name: str
    doc: str
    status: str
    rules: list[str] = field(default_factory=list)
    apis: list[tuple[str, str]] = field(default_factory=list)
    class_mentions: set[str] = field(default_factory=set)

    @property
    def domain(self) -> str:
        return "Master" if self.code.startswith("MASTER") else "Tenant"


def norm_path(path: str) -> str:
    path = path.split("?")[0].strip().rstrip("/")
    path = re.sub(r"\{[^}]*\}", "{}", path)
    return re.sub(r"/:\w+", "/{}", path)


def clean_md(text: str) -> str:
    return re.sub(r"\*\*|`", "", text).strip()


def parse_feature(md: Path) -> Feature:
    text = md.read_text(encoding="utf-8")
    lines = text.splitlines()
    title = next((l[2:].strip() for l in lines if l.startswith("# ")), md.stem)
    code_m = re.search(r"\*\*Code ID:\*\*\s*`([^`]+)`", text)
    status_m = re.search(r"\*\*Trạng thái:\*\*\s*`([^`]+)`", text)
    feat = Feature(
        code=code_m.group(1) if code_m else md.stem,
        name=title,
        doc=md.relative_to(ROOT).as_posix(),
        status=status_m.group(1) if status_m else "",
    )

    in_rules = False
    for line in lines:
        heading = re.match(r"^(#{1,6})\s+(.*)", line)
        if heading:
            in_rules = "business rule" in heading.group(2).lower()
            continue
        if not in_rules or not line.strip():
            continue
        item = re.match(r"^(\s*)(?:[-*]|\d+\.)\s+(.*)", line)
        if item and not item.group(1):
            feat.rules.append(clean_md(item.group(2)))
        elif feat.rules:
            feat.rules[-1] += " / " + clean_md(item.group(2) if item else line)

    seen: set[tuple[str, str]] = set()
    base_m = re.search(r"dưới\s+`(/api/[^`]+)`", text)
    base = base_m.group(1).rstrip("/") if base_m else ""

    def add_api(methods: str, path: str) -> None:
        if base and not path.startswith("/api"):
            path = base + ("" if path.startswith("?") or path == "/" else path)
        if not path.startswith("/") or path.startswith("/v3/"):
            return
        if not path.startswith("/api"):
            path = "/api/v1" + path
        optional = re.match(r"(.*)\[(.*)\]$", path)
        variants = [optional.group(1), optional.group(1) + optional.group(2)] if optional else [path]
        for method in re.split(r"[/,\s]+", methods.upper()):
            for variant in variants:
                key = (method, norm_path(variant))
                if method in HTTP and key not in seen:
                    seen.add(key)
                    feat.apis.append(key)

    for row in re.finditer(r"^\|\s*([A-Za-z/ ,]+?)\s*\|\s*`?([^`|]+?)`?\s*\|", text, re.M):
        add_api(row.group(1), row.group(2).strip())
    for inline in re.finditer(r"`(GET|POST|PUT|PATCH|DELETE)\s+(/[^`\s]+)`", text):
        add_api(inline.group(1), inline.group(2))

    feat.class_mentions = set(re.findall(r"`([A-Z][A-Za-z0-9]+)`", text))
    return feat


def load_features() -> tuple[list[Feature], list[str]]:
    """Docs without a Code ID (guides) are skipped; duplicated Code IDs get the doc name appended."""
    feats, skipped, used = [], [], set()
    for md in sorted(DOCS.rglob("*.md")):
        if md.name == "README.md":
            continue
        feat = parse_feature(md)
        if feat.code == md.stem:
            skipped.append(feat.doc)
            continue
        if feat.code in used:
            skipped.append(f"{feat.doc} (Code ID {feat.code} trùng, đổi thành {feat.code}-{md.stem})")
            feat.code = f"{feat.code}-{md.stem}"
        used.add(feat.code)
        feats.append(feat)
    return sorted(feats, key=lambda f: (f.domain != "Master", f.code)), skipped


# ------------------------------------------------------------- main code --
@dataclass
class Endpoint:
    method: str
    path: str
    controller: str
    handler: str


def string_literals(args: str) -> list[str]:
    return re.findall(r'"([^"]*)"', args or "")


def parse_controllers(classes: dict[str, Path]) -> tuple[list[Endpoint], dict[str, set[str]]]:
    endpoints: list[Endpoint] = []
    deps: dict[str, set[str]] = {}
    for name, path in classes.items():
        if not name.endswith("Controller"):
            continue
        src = path.read_text(encoding="utf-8")
        class_pos = re.search(r"\bclass\s+" + name, src)
        if not class_pos:
            continue
        header = src[: class_pos.start()]
        prefix_m = re.search(r"@RequestMapping\s*\(([^)]*)\)", header)
        prefixes = string_literals(prefix_m.group(1)) if prefix_m else [""]
        deps[name] = set(re.findall(r"private\s+final\s+(\w+)\s+\w+\s*;", src))
        body = src[class_pos.end():]
        for m in re.finditer(r"@(Get|Post|Put|Patch|Delete)Mapping(?:\s*\(([^)]*)\))?", body):
            handler_m = re.search(r"\s(\w+)\s*\(", body[m.end(): m.end() + 600])
            handler = handler_m.group(1) if handler_m else "?"
            for prefix in prefixes:
                for sub in string_literals(m.group(2)) or [""]:
                    endpoints.append(Endpoint(m.group(1).upper(), norm_path(prefix + sub), name, handler))
    return endpoints, deps


PUBLIC_METHOD = re.compile(
    r"^\s*public\s+(?!class|interface|record|enum|static\s+class)(?:static\s+)?(?:final\s+)?"
    r"[\w<>\[\],.?\s]+?\s+(\w+)\s*\(", re.M)


def public_methods(path: Path, cls: str) -> list[str]:
    names = PUBLIC_METHOD.findall(path.read_text(encoding="utf-8"))
    return list(dict.fromkeys(n for n in names if n != cls))


# ------------------------------------------------------------------ tests --
@dataclass
class TestCase:
    file: str
    cls: str
    method: str
    display: str
    subjects: list[str]
    kind: str


def guess_kind(name: str) -> str:
    low = name.lower()
    if any(w in low for w in ABNORMAL_WORDS):
        return "A"
    if any(w in low for w in BOUNDARY_WORDS):
        return "B"
    return "N"


def parse_tests(classes: dict[str, Path]) -> tuple[list[TestCase], dict[str, str]]:
    cases: list[TestCase] = []
    sources: dict[str, str] = {}
    for path in sorted(TEST.rglob("*.java")):
        src = path.read_text(encoding="utf-8")
        cls = path.stem
        sources[cls] = src
        subjects: list[str] = []
        base = re.sub(r"(Integration)?Tests?$", "", cls)
        if base in classes:
            subjects.append(base)
        subjects += re.findall(r"@InjectMocks\s+(?:private\s+)?(\w+)", src)
        subjects += re.findall(r"@WebMvcTest\s*\(\s*(?:controllers\s*=\s*)?\{?\s*(\w+)\.class", src)
        subjects += [c for c in re.findall(r"new\s+(\w+)\s*\(", src) if c.endswith(SUBJECT_SUFFIXES)]
        if not subjects:
            subjects = [c for c in set(re.findall(r"\b([A-Z]\w+)\b", src))
                        if c.endswith(SUBJECT_SUFFIXES) and c in classes]
        subjects = [s for s in dict.fromkeys(subjects) if s in classes]

        for m in re.finditer(r"@(?:Test|ParameterizedTest)\b", src):
            tail = src[m.end(): m.end() + 800]
            name_m = re.search(r"void\s+(\w+)\s*\(", tail)
            if not name_m:
                continue
            display_m = re.search(r'@DisplayName\s*\(\s*"([^"]*)"', tail[: name_m.start()])
            method = name_m.group(1)
            display = display_m.group(1) if display_m else ""
            cases.append(TestCase(path.relative_to(TEST).as_posix(), cls, method, display,
                                  subjects, guess_kind(method + " " + display)))
    return cases, sources


def load_jacoco() -> dict[str, tuple[int, int, int, int]]:
    """class -> (line_missed, line_covered, branch_missed, branch_covered)."""
    if not JACOCO_CSV.exists():
        return {}
    data: dict[str, tuple[int, int, int, int]] = {}
    with JACOCO_CSV.open(encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            cls = row["CLASS"].split(".")[-1]
            data[cls] = (int(row["LINE_MISSED"]), int(row["LINE_COVERED"]),
                         int(row["BRANCH_MISSED"]), int(row["BRANCH_COVERED"]))
    return data


def pct(missed: int, covered: int) -> str:
    total = missed + covered
    return f"{covered * 100 / total:.1f}%" if total else "n/a"


# ---------------------------------------------------------------- workbook --
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
GAP_FILL = PatternFill("solid", fgColor="FCE4D6")


def add_sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: list[int]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for row in rows:
        ws.append(row)
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=str(ROOT / ".tmp" / "testing" / "Traceability_Matrix.xlsx"))
    out = Path(parser.parse_args().out)

    classes = {p.stem: p for p in MAIN.rglob("*.java")}
    features, skipped_docs = load_features()
    endpoints, controller_deps = parse_controllers(classes)
    tests, test_sources = parse_tests(classes)
    jacoco = load_jacoco()

    # endpoint -> features, class -> features
    ep_index: dict[tuple[str, str], list[Endpoint]] = defaultdict(list)
    for ep in endpoints:
        ep_index[(ep.method, ep.path)].append(ep)

    def find_endpoints(method: str, path: str) -> list[Endpoint]:
        for candidate in (path, "/api/v1" + path):
            if (method, candidate) in ep_index:
                return ep_index[(method, candidate)]
        return []

    doc_only_apis: list[tuple[str, str, str]] = []
    for f in features:
        controllers = MASTER_CONTROLLERS.get(f.code)
        if controllers:
            doc_only_apis += [(f.code, m, p) for m, p in f.apis if not find_endpoints(m, p)]
            f.apis = sorted({(e.method, e.path) for e in endpoints if e.controller in controllers},
                            key=lambda mp: (mp[1], mp[0]))

    class_features: dict[str, set[str]] = defaultdict(set)
    documented: set[tuple[str, str]] = set()
    for f in features:
        for method, path in f.apis:
            for ep in find_endpoints(method, path):
                documented.add((ep.method, ep.path))
                class_features[ep.controller].add(f.code)
                for dep in controller_deps.get(ep.controller, ()):
                    class_features[dep].add(f.code)
        for mention in f.class_mentions & classes.keys():
            class_features[mention].add(f.code)

    feature_tests: dict[str, list[TestCase]] = defaultdict(list)
    for t in tests:
        for code in sorted({c for s in t.subjects for c in class_features.get(s, ())}):
            feature_tests[code].append(t)

    # ---- sheets
    wb = Workbook()
    wb.remove(wb.active)

    summary, rule_rows, api_rows = [], [], []
    for f in features:
        seq = 0

        def next_tc() -> str:
            nonlocal seq
            seq += 1
            return f"{f.code}-TC{seq:02d}"

        for i, rule in enumerate(f.rules, 1):
            rule_rows.append([f.code, f"{f.code}-BR{i:02d}", rule, next_tc(), guess_kind(rule), ""])
        for method, path in f.apis:
            eps = find_endpoints(method, path)
            impl = ", ".join(sorted({f"{e.controller}.{e.handler}" for e in eps})) or "KHÔNG TÌM THẤY"
            success_tc = next_tc()
            validation_tc = "" if method == "GET" else next_tc()
            api_rows.append([f.code, method, path, impl, success_tc, validation_tc, "", ""])
        if f.apis:
            api_rows.append([f.code, "*", "(chung cho feature)", "Security filter / TenantWebInterceptor",
                             "", "", next_tc(), next_tc()])
        matched = sum(1 for m, p in f.apis if find_endpoints(m, p))
        summary.append([f.code, f.name, f.domain, f.status, len(f.rules), len(f.apis), matched,
                        len(feature_tests.get(f.code, [])), seq, f.doc])

    add_sheet(wb, "1_Tong quan",
              ["Code", "Feature", "Domain", "Status doc", "# Business rule", "# API trong doc",
               "# API khớp code", "# @Test liên quan", "# TC dự kiến", "Tài liệu"],
              summary, [12, 45, 9, 12, 10, 10, 10, 10, 10, 55])
    add_sheet(wb, "2_Business Rules",
              ["Code", "BR ID", "Business rule", "TC ID", "Loại gợi ý (N/A/B)", "Ghi chú duyệt"],
              rule_rows, [12, 16, 90, 16, 10, 30])
    add_sheet(wb, "3_API Matrix",
              ["Code", "Method", "Path", "Controller.handler", "TC thành công", "TC lỗi validate",
               "TC 401/403", "TC sai tenant"],
              api_rows, [12, 8, 50, 45, 16, 16, 16, 16])

    test_rows = []
    for t in tests:
        codes = sorted({c for s in t.subjects for c in class_features.get(s, ())})
        test_rows.append([t.file, t.method, t.display, t.kind, ", ".join(t.subjects),
                          ", ".join(codes) or "CHƯA ÁNH XẠ"])
    add_sheet(wb, "4_Existing Tests",
              ["File test", "@Test method", "@DisplayName", "Loại gợi ý (N/A/B)", "Class được test",
               "Feature"], test_rows, [55, 50, 40, 10, 40, 25])

    service_rows = []
    for name, path in sorted(classes.items()):
        if not name.endswith("Service") or "/test/" in path.as_posix():
            continue
        if re.search(r"\binterface\s+" + name, path.read_text(encoding="utf-8")):
            continue
        refs = [src for src in test_sources.values() if re.search(r"\b" + name + r"\b", src)]
        cov = jacoco.get(name)
        for method in public_methods(path, name):
            calls = sum(len(re.findall(r"\." + method + r"\s*\(", src)) for src in refs)
            service_rows.append([
                ", ".join(sorted(class_features.get(name, ()))), name, method, calls,
                "Có" if calls else "CHƯA TEST",
                pct(cov[0], cov[1]) if cov else "", pct(cov[2], cov[3]) if cov else "",
            ])
    add_sheet(wb, "5_Service Methods",
              ["Feature", "Service", "Public method", "# lần gọi trong test", "Đã có test",
               "Line cov (class)", "Branch cov (class)"],
              service_rows, [20, 40, 40, 12, 12, 12, 12])

    gaps = [["Tài liệu bị bỏ qua / Code ID trùng", "", doc] for doc in skipped_docs]
    for row in summary:
        if row[7] == 0:
            gaps.append(["Feature chưa có @Test nào", row[0], row[1]])
        if row[4] == 0:
            gaps.append(["Feature không có business rule trong doc", row[0], row[1]])
    for row in api_rows:
        if row[3] == "KHÔNG TÌM THẤY":
            gaps.append(["API trong doc không có trong code", row[0], f"{row[1]} {row[2]}"])
    for code, method, path in doc_only_apis:
        gaps.append(["API trong doc Master lệch code (test theo code)", code, f"{method} {path}"])
    for (method, path), eps in sorted(ep_index.items(), key=lambda kv: kv[0][1]):
        if (method, path) not in documented:
            gaps.append(["Endpoint trong code chưa có trong doc", eps[0].controller, f"{method} {path}"])
    for row in test_rows:
        if row[5] == "CHƯA ÁNH XẠ":
            gaps.append(["@Test chưa ánh xạ được feature", row[0], row[1]])
    add_sheet(wb, "6_Gaps", ["Loại thiếu sót", "Code / Class", "Chi tiết"], gaps, [40, 40, 80])
    for row in wb["6_Gaps"].iter_rows(min_row=2):
        for cell in row:
            cell.fill = GAP_FILL

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    total_tc = sum(r[8] for r in summary)
    print(f"Features: {len(features)} | Business rules: {len(rule_rows)} | Doc APIs: "
          f"{sum(r[5] for r in summary)} (khớp code {sum(r[6] for r in summary)}) | "
          f"Endpoints code: {len(ep_index)}")
    print(f"@Test: {len(tests)} | TC dự kiến (Test Report): {total_tc} | "
          f"Service methods: {len(service_rows)} (chưa test {sum(r[4] == 'CHƯA TEST' for r in service_rows)})")
    print(f"JaCoCo csv: {'có' if jacoco else 'chưa có (chạy mvnw test jacoco:report)'} | Gaps: {len(gaps)}")
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
